#!/usr/bin/env python3
"""Прочитать выгрузку prolong.xlsx и разложить её в JSON для страниц отчёта.

Выгрузка — единственный источник цифр. Строка = report_dt × product ×
communication × scenario_group × scenario, в колонках метрики дерева и отток в двух разрезах:
_rur (рубли) и _cnt (клиенты). Скрипт не считает бизнес-логику: переносит
значения, достраивает дельты м/м, суммирует сценарии в группы и описывает
структуру дерева, воронок и коммуникаций — в выгрузке лежат только цифры.

Какие строки куда идут:
  product = deposit/casco/osago, communication = ALL, scenario = ALL
                                          → продуктовая воронка;
  product = deposit, communication = ALL, scenario = ALL
                                          → дерево и портфель;
  ALL / promo / no_promo, scenario = ALL → три воронки;
  900, scenario = ALL                    → покрытие сервисными коммуникациями;
  все коммуникации × сценарии, outflow   → блок «Сценарии оттока».

На выходе два файла:
  outflow-tree-graph.json  — дерево оттоков, узлы с рядами по обоим разрезам;
  prolongation-data.json   — портфель, воронки, сценарии и бенчмарк.

Бенчмарк рынка берётся с листа «Бенчмарк» той же книги, а если его нет —
из outflow-report-metrics.xlsx.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "prolong.xlsx"
DEFAULT_BENCHMARK = ROOT / "outflow-report-metrics.xlsx"
DEFAULT_TREE = ROOT / "outflow-tree-graph.json"
DEFAULT_PAGE = ROOT / "prolongation-data.json"

BILLION = 1e9
MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
             "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

MEASURES = [
    {"id": "rur", "label": "Деньги", "unit": "млрд ₽", "short": "₽"},
    {"id": "clnt", "label": "Клиенты", "unit": "клиентов", "short": "клиентов"},
]
# Разрез страницы → суффиксы колонки выгрузки, первый — основной.
SUFFIX = {"rur": ("rur",), "clnt": ("cnt", "clnt")}
DEPOSIT_PRODUCT = "deposit"
PRODUCT_LABELS = {"deposit": "Вклады", "casco": "КАСКО", "osago": "ОСАГО"}

# В новых выгрузках снижение остатка хранится отдельной метрикой. Старое
# имя оставляем запасным вариантом, чтобы уже сохранённые книги продолжали
# собираться до обновления Excel-файла.
METRIC_ALIASES = {
    "balance_decrease": ("balance_decrease", "tree_portfolio_prolong_erosion"),
}


# ── структура дерева ──────────────────────────────────────────────────────
# Книга хранит цифры, а форму дерева — этот список: кто чей родитель, в какой
# полосе стоит и где именно. Координаты подобраны так, чтобы кружки и подписи
# не налезали друг на друга.
TREE = [
    {"id": "portfolio", "metric": "tree_portfolio",
     "label": "Портфель", "kind": "hero", "goodWhen": "up",
     "level": 0, "parent": None, "x": 0, "y": -350},
    {"id": "prolong", "metric": "tree_portfolio_prolong",
     "label": "Ожидается пролонгация", "kind": "driver", "goodWhen": "up",
     "level": 1, "parent": "portfolio", "x": -260, "y": -205, "impact": "neutral"},
    {"id": "not_prolong_erosion", "metric": "tree_portfolio_not_prolong_erosion",
     "label": "Вывели до пролонгации", "kind": "outflow", "goodWhen": "down",
     "level": 1, "parent": "portfolio", "x": 300, "y": -205, "impact": "negative"},
    {"id": "prolong_passed", "metric": "tree_portfolio_prolong_passed",
     "label": "Прошли пролонгацию", "kind": "inflow", "goodWhen": "up",
     "level": 2, "parent": "prolong", "x": -400, "y": -60, "impact": "positive"},
    {"id": "prolong_erosion", "metric": "tree_portfolio_prolong_erosion",
     "label": "Ушли на пролонгации", "kind": "outflow", "goodWhen": "down",
     "level": 2, "parent": "prolong", "x": -60, "y": -60, "impact": "negative"},
    {"id": "settlement", "metric": "tree_portfolio_prolong_settlement",
     "label": "Осело", "kind": "inflow", "goodWhen": "up",
     "level": 3, "parent": "prolong_passed", "x": -520, "y": 85, "impact": "positive"},
    {"id": "left_after_30d", "metric": "tree_portfolio_prolong_left_after_30d",
     "label": "Вывели за 30 дней", "kind": "outflow", "goodWhen": "down",
     "level": 3, "parent": "prolong_passed", "x": -200, "y": 85, "impact": "negative"},
]


BANDS = [
    {"short": "Портфель", "full": "Портфель"},
    {"short": "Подход к дате", "full": "Подход к дате пролонгации"},
    {"short": "Исход", "full": "Исход пролонгации"},
    {"short": "Результат 30 дней", "full": "Результат за 30 дней"},
]

# Отток портфеля приходит готовой колонкой: это своя метрика витрины —
# клиент, потерявший ≥50% баланса за месяц. Сумме веток дерева она не равна
# (ветки считаются по когорте пролонгации), поэтому складывать их нельзя.
OUTFLOW_METRIC = "outflow"

TOTAL = "ALL"
NO_OUTFLOW = "no_outflow"
SERVICE = "900"

# ── воронки ───────────────────────────────────────────────────────────────
# Воронка = срез коммуникации, шаг = колонка дерева. level — ступень,
# leak — плашка потерь, gain — плашка удержанного.
FUNNEL_STEPS = [
    (2, "expected", "tree_portfolio_prolong", "level"),
    (3, "prolonged", "tree_portfolio_prolong_passed", "level"),
    (4, "left30", "tree_portfolio_prolong_left_after_30d", "level"),
    (5, "erosion30", "balance_decrease", "level"),
    (6, "net_outflow", "outflow", "level"),
    (7, "not_waited", "tree_portfolio_not_prolong_erosion", "leak"),
    (8, "settled", "tree_portfolio_prolong_settlement", "gain"),
]
FUNNELS = [
    {"id": "all", "communication": "ALL", "name": "Отток общий", "note": "Вся база",
     "names": {"prolonged": "Пролонгация прошла"}},
    {"id": "promo", "communication": "promo", "name": "Отток с промо",
     "note": "Когорта с промо-коммуникацией",
     "names": {"prolonged": "Получили Промо-коммуникацию"}},
    {"id": "base", "communication": "no_promo", "name": "Отток обычный",
     "note": "Когорта без промо-коммуникации",
     "names": {"prolonged": "Не получили Промо-коммуникацию"}},
]
STEP_NAMES = {
    "expected": "Ожидается пролонгация",
    "left30": "Закрыли в окне 30 дней",
    "erosion30": "Вывели со сбережений",
    "net_outflow": "Вывели со сбережений >50%",
    "not_waited": "Вывели до пролонгации",
    "settled": "Осело через 30 дней",
}

# ── сценарии и коммуникации ───────────────────────────────────────────────
# Подписи для страницы. Сценарий, которого здесь нет, выйдет под своим id.
SCENARIO_GROUPS = {
    "transfers_atm": "Снятие в банкомате",
    "transfers_m2m": "Переводы M2M",
    "transfers_p2p": "Переводы P2P",
    "saving": "Перевод в накопления",
    "PAYMENTS": "Платежи",
    "bal_dc": "Оставили на ДК",
    "other": "Прочее",
}


def scenario_group(group: str, scenario: str) -> str:
    """Разделяет старую общую группу переводов на актуальные группы."""
    if group != "transfers_atm":
        return group
    if scenario == "m2m":
        return "transfers_m2m"
    if scenario == "p2p" or ("p2p" in scenario and "atm" not in scenario):
        return "transfers_p2p"
    return group
SCENARIO_NAMES = {
    "payments": "Платежи",
    "pos": "POS-оборот",
    "big_payments": "Крупные платежи",
    "pos_big_payments": "POS + крупные платежи",
    "oms": "ОМС",
    "invest": "Инвестиции",
    "invest_oms": "Инвестиции + ОМС",
    "housing_loan": "Ипотека",
    "auto_loan": "Автокредит",
    "pl_loan": "Потребительский кредит",
    "pl_new": "Потребительский кредит",
    "housing_loan_new": "Новая ипотека",
    "auto_loan_new": "Новый автокредит",
    "pl_loan_new": "Новый потребительский кредит",
    "atm": "Снятие в банкомате",
    "m2m": "M2M в другие банки",
    "p2p": "P2P-переводы",
    "p2p_sber": "P2P внутри Сбера",
    "p2p_otherbank": "P2P в другой банк",
    "m2m_p2p_external": "M2M + P2P вовне",
    "m2m_p2p_external_atm": "M2M + P2P вовне + банкомат",
    "other": "Прочее",
    "external": "Внешние операции",
    "bal_dc": "Оставили на балансе Дебетовой карты",
}
# Все сценарии из выгрузки показываются внутри своих групп.
HIDDEN_SCENARIOS: set[str] = set()
# Коммуникации: parent задаёт уровень в санки. Сервисные (900) в санки не
# идут — у них своя карточка. Ветки промо пересекаются: клиент может
# получить несколько предложений, сумма детей больше родителя.
COMMUNICATIONS = [
    {"id": "900", "name": "Сервисные", "kind": "service", "parent": None},
    {"id": "promo", "name": "Промо", "kind": "promo", "parent": None},
    {"id": "no_promo", "name": "Без промо", "kind": "none", "parent": None},
    {"id": "stop_list", "name": "Стоп-лист", "kind": "stop", "parent": None},
    {"id": "promo.deposit", "name": "Вклад", "kind": "promo", "parent": "promo"},
    {"id": "promo.deposit.prolongation", "name": "Вклад · пролонгация", "kind": "promo", "parent": "promo"},
    {"id": "promo.saving", "name": "НС", "kind": "promo", "parent": "promo"},
    {"id": "promo.invest", "name": "Инвестиции", "kind": "promo", "parent": "promo"},
    {"id": "promo.oms", "name": "ОМС", "kind": "promo", "parent": "promo"},
    {"id": "promo.deposit.draft", "name": "Черновик вклада", "kind": "draft", "parent": "promo.deposit"},
    {"id": "promo.saving.draft", "name": "Черновик НС", "kind": "draft", "parent": "promo.saving"},
    {"id": "promo.invest.draft", "name": "Черновик инвестиций", "kind": "draft", "parent": "promo.invest"},
    {"id": "promo.oms.draft", "name": "Черновик ОМС", "kind": "draft", "parent": "promo.oms"},
]


def period_id(report_dt) -> str:
    """report_dt в любом виде Excel → «ГГГГ-ММ».

    Бывает датой, текстом «2026-06-30» или «30.06.2026» и числом — серийным
    номером дня Excel, если у колонки числовой формат.
    """
    if isinstance(report_dt, (datetime, date)):
        return report_dt.strftime("%Y-%m")
    if isinstance(report_dt, (int, float)):
        return (date(1899, 12, 30) + timedelta(days=int(report_dt))).strftime("%Y-%m")
    text = str(report_dt).strip()
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S", "%Y-%m"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m")
        except ValueError:
            pass
    raise SystemExit(f"Не разобрать report_dt: {report_dt!r}")


def key_text(value) -> str:
    """Ключ из ячейки: без пробелов по краям; 900.0 из числовой ячейки → «900».

    ALL и no_outflow сравниваем без учёта регистра — приводим к одному виду.
    """
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.upper() == TOTAL:
        return TOTAL
    if text.lower() == NO_OUTFLOW:
        return NO_OUTFLOW
    return text


def product_key(value) -> str:
    """Нормализовать product, сохранив старые книги без этой колонки."""
    if value is None or str(value).strip() == "":
        return DEPOSIT_PRODUCT
    text = key_text(value).lower()
    return {"deposits": DEPOSIT_PRODUCT, "вклад": DEPOSIT_PRODUCT,
            "вклады": DEPOSIT_PRODUCT}.get(text, text)


def number(value) -> float | None:
    """Число из ячейки: float, int или текст в любой локали.

    Понимаем «1 234,5», «1234.5», «1.234,56», «219,480,011.22», узкие и
    неразрывные пробелы, апостроф и «−» вместо минуса. Если в тексте есть и
    точка, и запятая, дробная часть — после последнего из них; один и тот же
    знак несколько раз подряд — разделитель тысяч.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[\s\u00a0\u202f\u2009']", "", str(value)).replace("\u2212", "-")
    if not text or text in ("-", "—"):
        return None
    if "," in text and "." in text:
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        text = text.replace(thousands, "").replace(decimal, ".")
    elif text.count(",") > 1:
        text = text.replace(",", "")
    elif text.count(".") > 1:
        text = text.replace(".", "")
    else:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        raise SystemExit(f"Не число в выгрузке: {value!r}") from None


def period_label(report_dt: str) -> str:
    text = str(report_dt)
    return f"{MONTHS_RU[int(text[5:7]) - 1]} {text[:4]}"


def scale_of(measure: str) -> float:
    return BILLION if measure == "rur" else 1


def sheet_rows(sheet) -> list[dict]:
    """Строки листа словарями по заголовку; хвостовые заметки отбрасываем."""
    # Заголовки без пробелов по краям и в нижнем регистре: «Outflow_RUR » = outflow_rur.
    headers = [str(cell.value).strip().lower() if cell.value is not None else None for cell in sheet[1]]
    rows = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        if raw[0] is None or raw[1] is None:
            continue
        row: dict = {}
        for head, value in zip(headers, raw):
            # Дубль заголовка: берём первую непустую ячейку, а не последнюю.
            if head and (head not in row or row[head] in (None, "")):
                row[head] = value
        rows.append(row)
    return rows


def with_deltas(values: dict[str, float]) -> dict[str, dict]:
    """Ряд значений → ряд со значением и дельтой м/м; у первого месяца дельты нет."""
    series = {}
    previous = None
    for key in sorted(values):
        value = values[key]
        delta = None if previous in (None, 0) else (value / previous - 1) * 100
        series[key] = {"rawValue": value, "deltaPct": delta}
        previous = value
    return series


class Export:
    """Выгрузка, проиндексированная по (период, продукт, коммуникация, сценарий)."""

    def __init__(self, rows: list[dict]):
        if not rows:
            raise SystemExit("Выгрузка пуста")
        self.has_product = any("product" in row for row in rows)
        self.cells: dict[tuple[str, str, str, str], dict] = {}
        self.blank: set[tuple[str, str, str, str]] = set()
        self.filled: set[tuple[str, str, str]] = set()
        self.groups: dict[str, str] = {}
        for row in rows:
            missing = [c for c in ("report_dt", "communication", "scenario_group", "scenario") if c not in row]
            if missing:
                raise SystemExit(f"В выгрузке нет колонок: {', '.join(missing)}")
            key = (period_id(row["report_dt"]), product_key(row.get("product")),
                   key_text(row["communication"]), key_text(row["scenario"]))
            if key in self.cells:
                raise SystemExit(f"Дубль строки: {' / '.join(key)}")
            self.cells[key] = row
            # Строка scenario = ALL — итог коммуникации, а не сценарий
            # оттока. Её нельзя включать в разрез сценариев: иначе она
            # дублирует весь отток и ломает сумму веток.
            if key[1] == DEPOSIT_PRODUCT and key[2] != TOTAL and key[3] != TOTAL:
                self.groups[key[3]] = scenario_group(
                    key_text(row["scenario_group"]), key[3]
                )
        self.periods = sorted({key[0] for key in self.cells})
        self.products = sorted({key[1] for key in self.cells})
        self.communications = sorted({key[2] for key in self.cells})

    def communications_for_product(self, product: str = DEPOSIT_PRODUCT) -> list[str]:
        return sorted({key[2] for key in self.cells if key[1] == product})

    def value(self, period: str, communication: str, scenario: str, metric: str,
              measure: str, required: bool = True, product: str = DEPOSIT_PRODUCT) -> float | None:
        row = self.cells.get((period, product, communication, scenario))
        metric_names = METRIC_ALIASES.get(metric, (metric,))
        columns = [f"{name}_{suffix}" for name in metric_names for suffix in SUFFIX[measure]]
        column = next((c for c in columns if row is not None and c in row), columns[0])
        value = None if row is None else number(row.get(column))
        # Пустой итог (scenario = ALL) собираем из сценариев: они не пересекаются.
        if value is None and row is not None and scenario == TOTAL:
            parts = [number(r.get(column)) for (p, pr, c, s), r in self.cells.items()
                     if p == period and pr == product and c == communication and s != TOTAL]
            if any(v is not None for v in parts):
                self.filled.add((period, communication, column))
                value = sum(v or 0.0 for v in parts)
        # Остальные пустые ячейки — ноль: выгрузка оставляет пустым то, чего не было.
        if value is None and row is not None and column in row:
            self.blank.add((period, communication, scenario, column))
            value = 0.0
        if value is None:
            if required:
                raise SystemExit(self.explain(period, communication, scenario, column, product))
            return None
        return value / scale_of(measure)

    def explain(self, period: str, communication: str, scenario: str, column: str,
                product: str = DEPOSIT_PRODUCT) -> str:
        """Почему значения нет: нет строки, нет колонки или пустая ячейка."""
        where = (f"{period} / {communication} / {scenario}" if product == DEPOSIT_PRODUCT
                 else f"{period} / {product} / {communication} / {scenario}")
        row = self.cells.get((period, product, communication, scenario))
        if row is None:
            same = sorted({k[2] for k in self.cells if k[0] == period and k[1] == product and k[3] == scenario})
            return (f"Нет строки {where}. Периоды в выгрузке: {', '.join(self.periods)}; "
                    f"коммуникации со scenario = {scenario}: {', '.join(same) or '—'}")
        if column not in row:
            near = [c for c in row if c.split("_")[0] == column.split("_")[0]]
            return (f"Нет колонки {column} (строка {where}). Похожие колонки: "
                    f"{', '.join(near) or '—'}")
        return f"Нет значения {column} в строке {where}"

    def total_series(self, communication: str, metric: str, measure: str,
                    product: str = DEPOSIT_PRODUCT) -> dict[str, float]:
        return {p: self.value(p, communication, TOTAL, metric, measure, product=product)
                for p in self.periods}


def warn_unsaved_formulas(workbook: Path, sheet) -> None:
    """Формула без сохранённого значения читается пустой — предупредить.

    Так бывает, если книгу собрал скрипт, а Excel её не пересчитал и не
    сохранил: в Excel цифра видна, а в файле её нет.
    """
    formulas = load_workbook(workbook, data_only=False)[sheet.title]
    lost = 0
    for with_values, with_formulas in zip(sheet.iter_rows(values_only=True),
                                          formulas.iter_rows(values_only=True)):
        lost += sum(1 for v, f in zip(with_values, with_formulas)
                    if v is None and isinstance(f, str) and f.startswith("="))
    if lost:
        print(f"  ВНИМАНИЕ: {lost} формул без сохранённого значения — они прочитаются пустыми. "
              "Откройте книгу в Excel, нажмите «Сохранить» и соберите заново.")


def read_export(workbook: Path) -> tuple[Export, bool, object]:
    if not workbook.is_file():
        raise SystemExit(f"Не найдена выгрузка: {workbook}")
    book = load_workbook(workbook, data_only=True)
    sheet = book["prolong"] if "prolong" in book.sheetnames else book.worksheets[0]
    warn_unsaved_formulas(workbook, sheet)
    synthetic = False
    if "Справка" in book.sheetnames:
        about = {row[0]: row[1] for row in book["Справка"].iter_rows(values_only=True) if row and row[0]}
        synthetic = str(about.get("synthetic", "")).lower() in ("yes", "true", "1")
    return Export(sheet_rows(sheet)), synthetic, book


def build_tree(export: Export) -> dict:
    nodes = []
    for spec in TREE:
        nodes.append({
            "id": spec["id"],
            "label": spec["label"],
            "kind": spec["kind"],
            "level": spec["level"],
            "lane": spec["x"],
            "goodWhen": spec["goodWhen"],
            "planned": False,
            "synthetic": False,
            "x": spec["x"],
            "y": spec["y"],
            "series": {m["id"]: with_deltas(export.total_series(TOTAL, spec["metric"], m["id"]))
                       for m in MEASURES},
        })

    links = []
    for spec in TREE:
        if not spec["parent"]:
            continue
        links.append({
            "source": spec["id"],
            "target": spec["parent"],
            "impact": spec["impact"],
            "relation": None,
            "label": "",          # подпись собирает страница из активного разреза
            "analyticsId": None,
        })

    return {
        "meta": {
            "source": "prolong.xlsx, строки communication = ALL, scenario = ALL",
            "periods": [{"id": p, "label": period_label(p + "-01")} for p in reversed(export.periods)],
            "measures": MEASURES,
            "bands": BANDS,
        },
        "nodes": nodes,
        "links": links,
    }


def build_portfolio(export: Export) -> dict:
    """Портфель и суммарный отток — для страницы воронки."""
    return {
        m["id"]: {
            "total": with_deltas(export.total_series(TOTAL, "tree_portfolio", m["id"])),
            "outflow": with_deltas(export.total_series(TOTAL, OUTFLOW_METRIC, m["id"])),
        }
        for m in MEASURES
    }


def build_products(export: Export) -> list[dict]:
    """Продуктовые ряды из communication = ALL, scenario = ALL."""
    products = []
    for product in export.products:
        values = {}
        for period in export.periods:
            values[period] = {
                measure["id"]: {
                    "total": export.value(period, TOTAL, TOTAL, "tree_portfolio", measure["id"],
                                           product=product),
                    "expected": export.value(period, TOTAL, TOTAL, "tree_portfolio_prolong", measure["id"],
                                              product=product),
                    "prolonged": export.value(period, TOTAL, TOTAL, "tree_portfolio_prolong_passed", measure["id"],
                                               product=product),
                    "outflow": export.value(period, TOTAL, TOTAL, OUTFLOW_METRIC, measure["id"],
                                             product=product),
                }
                for measure in MEASURES
            }
        products.append({
            "id": product,
            "label": PRODUCT_LABELS.get(product, product),
            "values": values,
            "communication": TOTAL,
        })
    return products


def build_funnels(export: Export) -> list[dict]:
    funnels = []
    deposit_communications = export.communications_for_product(DEPOSIT_PRODUCT)
    for spec in FUNNELS:
        if spec["communication"] not in deposit_communications:
            print(f"  внимание: в выгрузке нет коммуникации {spec['communication']}, "
                  f"воронки «{spec['name']}» не будет")
            continue
        steps = [{"no": no, "id": step, "name": spec["names"].get(step, STEP_NAMES.get(step, step)),
                  "kind": kind} for no, step, _, kind in FUNNEL_STEPS]
        steps.sort(key=lambda step: (step["kind"] != "level", step["no"]))
        series = {}
        for measure in MEASURES:
            series[measure["id"]] = {
                period: {step: export.value(period, spec["communication"], TOTAL, metric, measure["id"])
                         for _, step, metric, _ in FUNNEL_STEPS}
                for period in export.periods
            }
        funnels.append({"id": spec["id"], "name": spec["name"], "note": spec["note"],
                        "steps": steps, "series": series})
    return funnels


def build_service(export: Export) -> dict | None:
    """Покрытие базы пролонгации сервисными коммуникациями (900) по месяцам."""
    if SERVICE not in export.communications_for_product(DEPOSIT_PRODUCT):
        print(f"  внимание: в выгрузке нет коммуникации {SERVICE}, карточки сервисного покрытия не будет")
        return None
    by_period = {}
    for period in export.periods:
        # Коммуникация 900 относится только к депозитам. В продуктовых
        # строках КАСКО/ОСАГО такой записи может не быть — это не ошибка
        # выгрузки и не должно останавливать сборку.
        covered = export.value(period, SERVICE, TOTAL, "tree_portfolio_prolong", "rur",
                                required=False, product=DEPOSIT_PRODUCT)
        total = export.value(period, TOTAL, TOTAL, "tree_portfolio_prolong", "rur",
                             required=False, product=DEPOSIT_PRODUCT)
        if covered is None or total is None:
            continue
        by_period[period] = {
            m["id"]: {
                "covered": export.value(period, SERVICE, TOTAL, "tree_portfolio_prolong", m["id"],
                                         required=False, product=DEPOSIT_PRODUCT),
                "total": export.value(period, TOTAL, TOTAL, "tree_portfolio_prolong", m["id"],
                                       required=False, product=DEPOSIT_PRODUCT),
            }
            for m in MEASURES
        }
    if not by_period:
        return None
    latest = export.periods[-1]
    if latest not in by_period:
        latest = sorted(by_period)[-1]
    return {
        "period": latest,
        "source": "prolong.xlsx, communication = 900",
        "basis": "сумма портфеля и уникальные клиенты",
        "baseLabel": "Клиенты с пролонгацией в месяце",
        "note": "",
        "values": by_period[latest],
        "byPeriod": by_period,
    }


def build_scenarios(export: Export, synthetic: bool) -> dict:
    """Отток по сценариям × коммуникациям: сценарии суммируются в группы.

    Сценарии взаимоисключающие, поэтому группа = сумма своих сценариев, а
    сумма всех сценариев должна сойтись со строкой scenario = ALL.
    """
    all_scenario_ids = sorted(export.groups, key=lambda s: (export.groups[s], s))
    scenario_ids = sorted(
        (scenario for scenario in export.groups if scenario not in HIDDEN_SCENARIOS),
        key=lambda s: (export.groups[s], s),
    )
    group_ids = []
    for scenario in scenario_ids:
        group = export.groups[scenario]
        if group != NO_OUTFLOW and group not in group_ids:
            group_ids.append(group)

    deposit_communications = export.communications_for_product(DEPOSIT_PRODUCT)
    known = {c["id"] for c in COMMUNICATIONS}
    extra = [c for c in deposit_communications if c not in known and c != TOTAL]
    if extra:
        print(f"  внимание: коммуникации без описания, в сценарии не попали: {', '.join(extra)}")

    values = {}
    for period in export.periods:
        values[period] = {}
        for comm in [TOTAL] + [c["id"] for c in COMMUNICATIONS]:
            if comm not in deposit_communications:
                continue
            has_total = (period, DEPOSIT_PRODUCT, comm, TOTAL) in export.cells
            has_scenarios = any(
                p == period and product == DEPOSIT_PRODUCT and communication == comm and scenario != TOTAL
                for p, product, communication, scenario in export.cells
            )
            # Коммуникация может быть заведена только в части периодов
            # (например, 900). Если в конкретном месяце её нет, не создаём
            # пустой разрез и не требуем строку «... / <коммуникация> / ALL».
            if not has_total and not has_scenarios:
                continue
            block = {"total": {}, "groups": {}, "scenarios": {}}
            for measure in MEASURES:
                mid = measure["id"]
                total = export.value(period, comm, TOTAL, OUTFLOW_METRIC, mid, required=False)
                summed = 0.0
                for scenario in all_scenario_ids:
                    group = export.groups[scenario]
                    if group == NO_OUTFLOW:
                        continue
                    value = export.value(period, comm, scenario, OUTFLOW_METRIC, mid, required=False) or 0.0
                    if scenario not in HIDDEN_SCENARIOS:
                        block["scenarios"].setdefault(scenario, {})[mid] = value
                    block["groups"].setdefault(group, {mid: 0.0 for mid in SUFFIX})[mid] += value
                    summed += value
                # Если строка scenario = ALL не передана, итог берём из
                # взаимоисключающих сценариев. Для коммуникационных разрезов
                # эта строка не обязательна.
                block["total"][mid] = summed if total is None else total
                if total and has_scenarios and abs(summed - total) / total > 0.005:
                    print(f"  внимание: {period} / {comm} / {mid}: сумма сценариев {summed:,.2f} "
                          f"не сходится с ALL {total:,.2f}")
            values[period][comm] = block

    return {
        "metric": "outflow",
        "synthetic": synthetic,
        "groups": [{"id": g, "name": SCENARIO_GROUPS.get(g, g),
                    "scenarios": [{"id": s, "name": SCENARIO_NAMES.get(s, s)}
                                  for s in scenario_ids if export.groups[s] == g]}
                   for g in group_ids],
        "communications": [c for c in COMMUNICATIONS if c["id"] in deposit_communications],
        "values": values,
    }


def read_benchmark(book, fallback: Path) -> list[dict]:
    if "Бенчмарк" in book.sheetnames:
        sheet = book["Бенчмарк"]
    elif fallback.is_file() and "Бенчмарк" in load_workbook(fallback, data_only=True).sheetnames:
        sheet = load_workbook(fallback, data_only=True)["Бенчмарк"]
    else:
        print("  внимание: листа «Бенчмарк» нет, блок AI Benchmark будет пустым")
        return []
    points = []
    for row in sheet_rows(sheet):
        points.append({
            "period": str(row["period"]),
            "measure": row["measure"],
            "metric": row["metric_id"],
            "valuePct": None if row.get("value_pct") is None else float(row["value_pct"]) * 100,
            "source": row.get("source_name") or "",
            "url": row.get("url") or "",
            "note": row.get("note") or "",
        })
    return points


def build(workbook: Path, tree_out: Path, page_out: Path, benchmark: Path) -> None:
    export, synthetic, book = read_export(workbook)

    graph = build_tree(export)
    funnels = build_funnels(export)
    scenarios = build_scenarios(export, synthetic)
    page = {
        "meta": {
            "source": workbook.name,
            "periods": graph["meta"]["periods"],
            "measures": MEASURES,
            "productDimension": export.has_product,
            "outflowDef": "колонка outflow выгрузки: клиент потерял ≥50% баланса за месяц",
        },
        "portfolio": build_portfolio(export),
        "products": build_products(export),
        "serviceCoverage": build_service(export),
        "funnels": funnels,
        "scenarios": scenarios,
        "benchmark": read_benchmark(book, benchmark),
    }

    if export.filled:
        columns = sorted({column for _, _, column in export.filled})
        print(f"  внимание: пустые итоги scenario = ALL ({len(export.filled)} шт.) собраны суммой "
              f"сценариев; колонки: {', '.join(columns)}")
    if export.blank:
        columns = sorted({column for *_, column in export.blank})
        print(f"  внимание: {len(export.blank)} пустых ячеек посчитаны нулём; "
              f"колонки: {', '.join(columns)}")

    for path, payload in [(tree_out, graph), (page_out, page)]:
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"{tree_out.name}: {len(graph['nodes'])} узлов, {len(export.periods)} периодов, "
          f"{len(MEASURES)} разреза")
    print(f"{page_out.name}: {len(funnels)} воронки, {len(scenarios['groups'])} групп сценариев, "
          f"{len(scenarios['communications'])} коммуникаций, {len(page['benchmark'])} строк бенчмарка"
          + (" · синтетика" if synthetic else ""))


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--benchmark", type=Path, default=DEFAULT_BENCHMARK,
                        help="книга с листом «Бенчмарк», если его нет в выгрузке")
    parser.add_argument("--tree-out", type=Path, default=DEFAULT_TREE)
    parser.add_argument("--page-out", type=Path, default=DEFAULT_PAGE)
    return parser.parse_args(argv)


def main() -> None:
    args = parse_args()
    build(args.workbook, args.tree_out, args.page_out, args.benchmark)


if __name__ == "__main__":
    sys.exit(main())
