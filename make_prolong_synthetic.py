#!/usr/bin/env python3
"""Собрать синтетический prolong.xlsx в формате боевой выгрузки.

Боевая таблица лежит на другой машине; здесь — её аналог для отрисовки.
Строка = report_dt × communication × scenario_group × scenario, в колонках
метрики дерева и отток в двух разрезах: _rur (рубли) и _cnt (клиенты).

Что настоящее: итоги ALL / promo / no_promo по всем сценариям (строки
scenario = ALL) взяты из листа «Дерево» outflow-report-metrics.xlsx.
Всё остальное — разбивка по сценариям, прочие коммуникации — синтетика:
пропорции правдоподобные, цифры выдуманные.

Правила, которые соблюдает синтетика (их же ждёт сборка):
  * сценарии взаимоисключающие: сумма по сценариям = строка ALL/ALL;
  * no_outflow — не отток, outflow у него ноль;
  * коммуникации пересекаются: клиент может получить несколько промо;
  * черновик (…draft) — подмножество своего промо, промо — подмножество ALL.
"""

from __future__ import annotations

import argparse
import random
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = ROOT / "outflow-report-metrics.xlsx"
DEFAULT_OUTPUT = ROOT / "prolong.xlsx"

REPORT_DT = "2026-06-30"
SEED = 20260630

METRICS = [
    "tree_portfolio",
    "tree_portfolio_prolong",
    "tree_portfolio_not_prolong_erosion",
    "tree_portfolio_prolong_erosion",
    "tree_portfolio_prolong_passed",
    "tree_portfolio_prolong_left_after_30d",
    "tree_portfolio_prolong_settlement",
    "outflow",
]
MEASURES = ["rur", "cnt"]
COLUMNS = ["report_dt", "communication", "scenario_group", "scenario"] + [
    f"{metric}_{measure}" for metric in METRICS for measure in MEASURES
]

# Сценарии оттока: группа → сценарии и их вес в оттоке (рубли, клиенты).
SCENARIOS = {
    "transfers_atm": {
        "m2m": (14, 8), "p2p": (5, 9), "atm": (3, 6),
        "m2m_p2p_external": (9, 6), "m2m_p2p_external_atm": (4, 4),
    },
    "PAYMENTS": {"pos": (4, 13), "big_payments": (6, 4), "pos_big_payments": (3, 5)},
    "saving": {"invest": (9, 5), "oms": (3, 3), "invest_oms": (3, 2)},
    "loans": {"housing_loan": (8, 5), "auto_loan": (3, 3), "pl_loan": (3, 6)},
    "loans_new": {"housing_loan_new": (4, 2), "auto_loan_new": (1, 1), "pl_loan_new": (1, 3)},
    "other": {"other": (3, 4), "external": (3, 2)},
}

# Доля метрики, которая приходится на клиентов без оттока (no_outflow).
NO_OUTFLOW_SHARE = {
    "tree_portfolio": 0.93,
    "tree_portfolio_prolong": 0.90,
    "tree_portfolio_not_prolong_erosion": 0.25,
    "tree_portfolio_prolong_erosion": 0.20,
    "tree_portfolio_prolong_passed": 0.91,
    "tree_portfolio_prolong_left_after_30d": 0.35,
    "tree_portfolio_prolong_settlement": 0.96,
    "outflow": 0.0,
}

# Коммуникация → (родитель, доля от родителя). Корень — ALL.
# None — подгоняем под итог: promo и no_promo под настоящий, 900 — под долю
# SERVICE_SHARE; остальные — доля от родителя.
COMMUNICATIONS = {
    "900": ("ALL", None),
    "promo": ("ALL", None),
    "no_promo": ("ALL", None),
    "stop_list": ("ALL", 0.05),
    "promo.deposit": ("promo", 0.62),
    "promo.deposit.prolongation": ("promo", 0.48),
    "promo.saving": ("promo", 0.34),
    "promo.invest": ("promo", 0.14),
    "promo.oms": ("promo", 0.07),
    "promo.deposit.draft": ("promo.deposit", 0.24),
    "promo.saving.draft": ("promo.saving", 0.18),
    "promo.invest.draft": ("promo.invest", 0.15),
    "promo.oms.draft": ("promo.oms", 0.12),
}


# Сервисное покрытие базы пролонгации из prolongation-communications.json:
# 4 163 364 из 4 762 491 клиентов, 4080 из 4430 млрд ₽.
SERVICE_SHARE = {"rur": 4080 / 4430, "cnt": 4163364 / 4762491}


def scenario_list() -> list[tuple[str, str]]:
    pairs = [("no_outflow", "no_outflow")]
    for group, items in SCENARIOS.items():
        pairs += [(group, name) for name in items]
    return pairs


def read_real_totals(source: Path) -> dict[str, dict[str, float]]:
    """Итоги дерева по коммуникациям ALL / promo / no_promo из книги метрик."""
    book = load_workbook(source, data_only=True)
    sheet = book["Дерево"]
    head = [cell.value for cell in sheet[1]]
    totals = {}
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        if raw[0] is None:
            continue
        row = dict(zip(head, raw))
        values = {}
        for metric in METRICS:
            values[f"{metric}_rur"] = float(row[f"{metric}_rur"])
            values[f"{metric}_cnt"] = float(row[f"{metric}_clnt"])
        totals[row["communication"]] = values
    missing = {"ALL", "promo", "no_promo"} - set(totals)
    if missing:
        raise SystemExit(f"В листе «Дерево» нет коммуникаций: {', '.join(sorted(missing))}")
    return totals


def base_split(total: dict[str, float]) -> dict[tuple[str, str], dict[str, float]]:
    """Разложить итог ALL по сценариям: веса по оттоку, no_outflow — остаток."""
    pairs = scenario_list()
    out = {pair: {} for pair in pairs}
    for measure_index, measure in enumerate(MEASURES):
        weights = {(g, s): w[measure_index] for g, items in SCENARIOS.items() for s, w in items.items()}
        weight_sum = sum(weights.values())
        for metric in METRICS:
            column = f"{metric}_{measure}"
            rest = total[column] * (1 - NO_OUTFLOW_SHARE[metric])
            out[("no_outflow", "no_outflow")][column] = total[column] * NO_OUTFLOW_SHARE[metric]
            for pair, weight in weights.items():
                out[pair][column] = rest * weight / weight_sum
    return out


def coverage(rng: random.Random, pairs, level: float) -> dict[tuple[str, str], float]:
    """Покрытие коммуникацией по сценариям: общий уровень ± разброс по группам."""
    group_bias = {g: rng.uniform(0.85, 1.12) for g in list(SCENARIOS) + ["no_outflow"]}
    return {pair: level * group_bias[pair[0]] * rng.uniform(0.95, 1.05) for pair in pairs}


def fit(parent_cells, cov, column: str, target: float) -> dict:
    """Разложить target по сценариям пропорционально покрытию, не выше родителя.

    Ячейку, упёршуюся в 99% родителя, фиксируем и досыпаем остаток в прочие.
    """
    cap = {pair: parent_cells[pair][column] * 0.99 for pair in cov}
    fixed: dict = {}
    for _ in range(10):
        free = [pair for pair in cov if pair not in fixed]
        raw = {pair: parent_cells[pair][column] * cov[pair] for pair in free}
        raw_sum = sum(raw.values())
        scale = (target - sum(fixed.values())) / raw_sum if raw_sum else 0
        over = [pair for pair in free if raw[pair] * scale > cap[pair]]
        if not over:
            break
        for pair in over:
            fixed[pair] = cap[pair]
    return {pair: fixed.get(pair, parent_cells[pair][column] * cov[pair] * scale) for pair in cov}


def fix_rounding(rows: list[dict], target: dict[str, float]) -> None:
    """Досыпать хвост округления в крупнейшую строку: итог = настоящей цифре."""
    for column in COLUMNS[4:]:
        want = round(target[column]) if column.endswith("_cnt") else round(target[column], 2)
        gap = want - sum(r[column] for r in rows)
        # no_outflow не трогаем в колонке оттока: у него отток строго ноль.
        pool = [r for r in rows if not (column.startswith("outflow") and r["scenario"] == "no_outflow")]
        biggest = max(pool, key=lambda r: r[column])
        biggest[column] = round(biggest[column] + gap, 2) if column.endswith("_rur") else biggest[column] + gap


def build_rows(real: dict[str, dict[str, float]], seed: int) -> list[dict]:
    rng = random.Random(seed)
    pairs = scenario_list()
    cells: dict[str, dict[tuple[str, str], dict[str, float]]] = {"ALL": base_split(real["ALL"])}

    for comm, (parent, share) in COMMUNICATIONS.items():
        parent_cells = cells[parent]
        cells[comm] = {pair: {} for pair in pairs}
        for column in COLUMNS[4:]:
            if share is None:
                if comm in real:
                    target = real[comm][column]
                else:
                    target = real["ALL"][column] * SERVICE_SHARE[column.rsplit("_", 1)[1]]
                level = target / real["ALL"][column] if real["ALL"][column] else 0
                cov = coverage(rng, pairs, level)
                for pair, value in fit(parent_cells, cov, column, target).items():
                    cells[comm][pair][column] = value
            else:
                cov = coverage(rng, pairs, share)
                for pair in pairs:
                    cells[comm][pair][column] = parent_cells[pair][column] * min(cov[pair], 0.97)

    rows = []
    for comm in ["ALL"] + list(COMMUNICATIONS):
        comm_cells = cells[comm]
        scenario_rows = []
        for group, scenario in pairs:
            row = {"report_dt": REPORT_DT, "communication": comm,
                   "scenario_group": group, "scenario": scenario}
            for column in COLUMNS[4:]:
                value = comm_cells[(group, scenario)][column]
                row[column] = round(value) if column.endswith("_cnt") else round(value, 2)
            scenario_rows.append(row)
        if comm in real:
            fix_rounding(scenario_rows, real[comm])
        total = {"report_dt": REPORT_DT, "communication": comm,
                 "scenario_group": "ALL", "scenario": "ALL"}
        for column in COLUMNS[4:]:
            total[column] = sum(r[column] for r in scenario_rows)
            if column.endswith("_rur"):
                total[column] = round(total[column], 2)
        rows.append(total)
        rows += scenario_rows
    return rows


def copy_benchmark(source: Path, book: Workbook) -> None:
    """Бенчмарк рынка переезжает в ту же книгу — сборке нужен один файл."""
    src = load_workbook(source)
    if "Бенчмарк" not in src.sheetnames:
        return
    sheet = book.create_sheet("Бенчмарк")
    for raw in src["Бенчмарк"].iter_rows(values_only=True):
        sheet.append(list(raw))
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def write(rows: list[dict], output: Path, source: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "prolong"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row[c] for c in COLUMNS])
    sheet.freeze_panes = "E2"

    copy_benchmark(source, book)

    about = book.create_sheet("Справка")
    about.append(["key", "value"])
    about.append(["synthetic", "yes"])
    about.append(["note", "Итоги ALL / promo / no_promo (scenario = ALL) — настоящие, июнь 2026. "
                          "Разбивка по сценариям и прочие коммуникации — синтетика. "
                          "В боевой выгрузке этого листа нет, и отчёт снимает пометку «синтетика»."])
    book.save(output)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--seed", type=int, default=SEED)
    return parser.parse_args(argv)


def main() -> None:
    args = parse_args()
    rows = build_rows(read_real_totals(args.source), args.seed)
    write(rows, args.output, args.source)
    print(f"{args.output.name}: {len(rows)} строк, {len(COMMUNICATIONS) + 1} коммуникаций, "
          f"{len(scenario_list())} сценариев")


if __name__ == "__main__":
    sys.exit(main())
