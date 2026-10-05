#!/usr/bin/env python3
"""Build the questionnaire seed data from flat_table.xlsx."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = ROOT / "flat_table.xlsx"
DEFAULT_OUTPUT = ROOT / "gravity-app" / "src" / "questionnaire-app" / "questionnaire-data.json"
SHEET_NAME = "Лист1"
REQUIRED_COLUMNS = {
    "metric_code",
    "metric_name",
    "Юнит",
    "трайб",
    "product",
    "тип",
    "факт",
    "макс балл",
    "metric_group",
    "metric_subgroup",
    "metric_footer",
    "is_visible",
    "flg",
}


def number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        parsed = float(str(value).strip().replace(",", "."))
        return parsed if math.isfinite(parsed) else None
    except (TypeError, ValueError):
        return None


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def enabled(value: Any) -> bool:
    return number(value) == 1


def build_payload(input_path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"В {input_path.name} нет листа {SHEET_NAME!r}")

    rows = workbook[SHEET_NAME].iter_rows(values_only=True)
    header = [text(value) for value in next(rows)]
    missing = sorted(REQUIRED_COLUMNS - set(header))
    if missing:
        raise ValueError(f"В {SHEET_NAME!r} нет колонок: {', '.join(missing)}")
    indexes = {name: header.index(name) for name in REQUIRED_COLUMNS}

    products: dict[str, dict[str, Any]] = {}
    for source_row_number, raw_row in enumerate(rows, start=2):
        row = list(raw_row) + [None] * max(0, len(header) - len(raw_row))
        score = number(row[indexes["факт"]])
        max_score = number(row[indexes["макс балл"]])
        if (
            score is None
            or max_score is None
            or max_score <= 0
            or score >= max_score
            or text(row[indexes["тип"]]).casefold() == "исключить"
            or not enabled(row[indexes["is_visible"]])
            or not enabled(row[indexes["flg"]])
        ):
            continue

        product_name = text(row[indexes["product"]])
        metric_name = text(row[indexes["metric_name"]])
        if not product_name or not metric_name:
            continue

        product = products.setdefault(
            product_name,
            {
                "id": product_name,
                "name": product_name,
                "unit": text(row[indexes["Юнит"]]),
                "tribe": text(row[indexes["трайб"]]),
                "entityType": text(row[indexes["тип"]]),
                "questions": [],
            },
        )
        identity = "\x1f".join(
            [
                product_name,
                text(row[indexes["metric_code"]]),
                text(row[indexes["metric_group"]]),
                metric_name,
                str(source_row_number),
            ]
        )
        question_id = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:16]
        product["questions"].append(
            {
                "id": f"metric-{question_id}",
                "sourceRow": source_row_number,
                "metricCode": text(row[indexes["metric_code"]]),
                "metricName": metric_name,
                "metricGroup": text(row[indexes["metric_group"]]),
                "metricSubgroup": text(row[indexes["metric_subgroup"]]),
                "metricFooter": text(row[indexes["metric_footer"]]),
                "score": score,
                "maxScore": max_score,
                "question": "",
                "placeholder": "Введите ответ",
                "answerType": "text",
                "required": True,
            }
        )

    ordered_products = sorted(products.values(), key=lambda item: item["name"].casefold())
    for product in ordered_products:
        product["questions"].sort(
            key=lambda item: (
                item["metricGroup"].casefold(),
                item["metricName"].casefold(),
                item["sourceRow"],
            )
        )
    return {
        "source": input_path.name,
        "sheet": SHEET_NAME,
        "filter": "тип != Исключить; is_visible = 1; flg = 1; факт и макс балл — конечные числа; макс балл > 0; факт < макс балл",
        "productCount": len(ordered_products),
        "questionCount": sum(len(product["questions"]) for product in ordered_products),
        "products": ordered_products,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    payload = build_payload(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"{args.output}: {payload['productCount']} products, "
        f"{payload['questionCount']} questions"
    )


if __name__ == "__main__":
    main()
