"""Сборка JSON отчёта из выгрузки prolong.xlsx."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

import build_report_data as brd
import make_prolong_synthetic as synth

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "outflow-report-metrics.xlsx"

pytestmark = pytest.mark.skipif(not SOURCE.is_file(), reason="нет outflow-report-metrics.xlsx")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("prolong")
    workbook = tmp / "prolong.xlsx"
    real = synth.read_real_totals(SOURCE)
    synth.write(synth.build_rows(real, synth.SEED), workbook, SOURCE)
    tree, page = tmp / "tree.json", tmp / "page.json"
    brd.build(workbook, tree, page, SOURCE)
    return real, json.loads(tree.read_text()), json.loads(page.read_text())


def test_funnels_take_tree_columns_of_their_communication(built):
    real, _, page = built
    funnels = {f["id"]: f for f in page["funnels"]}
    assert set(funnels) == {"all", "promo", "base"}
    june = funnels["promo"]["series"]["clnt"]["2026-06"]
    assert june["expected"] == real["promo"]["tree_portfolio_prolong_cnt"]
    assert june["net_outflow"] == real["promo"]["outflow_cnt"]
    assert funnels["base"]["series"]["clnt"]["2026-06"]["settled"] == \
        real["no_promo"]["tree_portfolio_prolong_settlement_cnt"]


def test_tree_and_portfolio_use_all_row(built):
    real, tree, page = built
    portfolio = next(n for n in tree["nodes"] if n["id"] == "portfolio")
    assert portfolio["series"]["rur"]["2026-06"]["rawValue"] == pytest.approx(
        real["ALL"]["tree_portfolio_rur"] / 1e9)
    assert page["portfolio"]["clnt"]["outflow"]["2026-06"]["rawValue"] == real["ALL"]["outflow_cnt"]


def test_scenario_groups_sum_to_total_and_skip_no_outflow(built):
    _, _, page = built
    scenarios = page["scenarios"]
    assert scenarios["synthetic"] is True
    assert "no_outflow" not in {g["id"] for g in scenarios["groups"]}
    assert all(scenario["id"] != "atm"
               for group in scenarios["groups"] for scenario in group["scenarios"])
    for comm, block in scenarios["values"]["2026-06"].items():
        for measure in ("rur", "clnt"):
            summed = sum(g[measure] for g in block["groups"].values())
            assert summed == pytest.approx(block["total"][measure], rel=1e-6), comm


def test_communications_are_nested(built):
    _, _, page = built
    values = page["scenarios"]["values"]["2026-06"]
    for comm in page["scenarios"]["communications"]:
        parent = comm["parent"] or "ALL"
        for group, cell in values[comm["id"]]["groups"].items():
            assert cell["clnt"] <= values[parent]["groups"][group]["clnt"], (comm["id"], group)


def test_service_coverage_from_900(built):
    _, _, page = built
    service = page["serviceCoverage"]["byPeriod"]["2026-06"]["clnt"]
    assert service["covered"] / service["total"] == pytest.approx(0.874, abs=0.001)


def test_duplicate_rows_are_rejected():
    row = {"report_dt": "2026-06-30", "communication": "ALL", "scenario_group": "ALL", "scenario": "ALL"}
    with pytest.raises(SystemExit, match="Дубль"):
        brd.Export([row, dict(row)])


def test_real_export_without_help_sheet_is_not_synthetic(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(synth.COLUMNS)
    sheet.append(["2026-06-30", "ALL", "ALL", "ALL"] + [1] * (len(synth.COLUMNS) - 4))
    path = tmp_path / "prolong.xlsx"
    book.save(path)
    _, synthetic, _ = brd.read_export(path)
    assert synthetic is False


def test_keys_and_numbers_tolerate_excel_formats():
    columns = [c.upper() + " " if i == 0 else c for i, c in enumerate(synth.COLUMNS)]
    head = [c.strip().lower() for c in columns]
    values = ["1 234,5"] + [2] * (len(synth.COLUMNS) - 5)
    rows = [dict(zip(head, [46203, 900.0, "all ", " All"] + values))]
    export = brd.Export(rows)
    assert export.periods == ["2026-06"]
    assert export.communications == ["900"]
    assert export.value("2026-06", "900", "ALL", "tree_portfolio", "rur") == pytest.approx(1234.5 / 1e9)


def test_missing_row_fails_but_blank_cell_is_zero():
    row = {"report_dt": "30.06.2026", "communication": "promo", "scenario_group": "ALL", "scenario": "ALL",
           "outflow_rur": None}
    export = brd.Export([row])
    with pytest.raises(SystemExit, match="Нет строки 2026-06 / ALL / ALL.*promo"):
        export.value("2026-06", "ALL", "ALL", "outflow", "rur")
    assert export.value("2026-06", "promo", "ALL", "outflow", "rur") == 0.0
    assert export.blank == {("2026-06", "promo", "ALL", "outflow_rur")}


def test_balance_decrease_prefers_dedicated_excel_columns():
    row = {
        "report_dt": "2026-06-30",
        "communication": "ALL",
        "scenario_group": "ALL",
        "scenario": "ALL",
        "balance_decrease_cnt": 362631,
        "balance_decrease_rur": 223578157626,
        "tree_portfolio_prolong_erosion_cnt": 39510,
        "tree_portfolio_prolong_erosion_rur": 31834114941.04,
    }
    export = brd.Export([row])
    assert export.value("2026-06", "ALL", "ALL", "balance_decrease", "clnt") == 362631
    assert export.value("2026-06", "ALL", "ALL", "balance_decrease", "rur") == pytest.approx(
        223578157626 / 1e9
    )


def test_blank_total_is_summed_from_scenarios():
    base = {"report_dt": "2026-06-30", "communication": "ALL", "scenario_group": "PAYMENTS"}
    rows = [
        {**base, "scenario_group": "ALL", "scenario": "ALL", "outflow_rur": None},
        {**base, "scenario": "pos", "outflow_rur": 3},
        {**base, "scenario": "big_payments", "outflow_rur": 4},
    ]
    export = brd.Export(rows)
    assert export.value("2026-06", "ALL", "ALL", "outflow", "rur") == pytest.approx(7 / 1e9)
    assert export.filled == {("2026-06", "ALL", "outflow_rur")}


def test_duplicate_header_keeps_filled_cell(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(["report_dt", "communication", "scenario_group", "scenario", "outflow_rur", "outflow_rur"])
    sheet.append(["2026-06-30", "ALL", "ALL", "ALL", 5, None])
    assert brd.sheet_rows(sheet)[0]["outflow_rur"] == 5


def test_unsaved_formula_is_reported(tmp_path, capsys):
    book = Workbook()
    sheet = book.active
    sheet.append(["report_dt", "communication", "outflow_rur"])
    sheet.append(["2026-06-30", "ALL", "=1+1"])
    path = tmp_path / "f.xlsx"
    book.save(path)
    loaded = brd.load_workbook(path, data_only=True).active
    brd.warn_unsaved_formulas(path, loaded)
    assert "1 формул без сохранённого значения" in capsys.readouterr().out


@pytest.mark.parametrize("raw, expected", [
    (219.48, 219.48), (12, 12.0), ("219,48", 219.48), ("219.48", 219.48),
    ("1 234,5", 1234.5), ("1 234,5", 1234.5), ("1 234,5", 1234.5),
    ("1.234,56", 1234.56), ("219,480,011.22", 219480011.22), ("1.234.567", 1234567.0),
    ("1,234,567", 1234567.0), ("−5,5", -5.5), ("2,1948E+11", 2.1948e11), ("", None), (None, None),
])
def test_number_reads_any_locale(raw, expected):
    assert brd.number(raw) == (pytest.approx(expected) if expected is not None else None)
